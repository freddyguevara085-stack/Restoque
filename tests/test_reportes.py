from io import BytesIO
import openpyxl


def test_dashboard_metrics_aggregation(client, admin_headers):
    # Seed a paca
    p_res = client.post("/api/pacas/", json={
        "descripcion": "Paca Reportes",
        "costo": 2000.0,
        "peso_lbs": 30.0,
        "categorias": [
            {"nombre": "Zapatos", "cantidad_total": 10, "precio_venta": 300.0}
        ]
    }, headers=admin_headers)
    cat_id = p_res.json()["categorias"][0]["id"]

    # Sell 2 items (2 * 300 = 600)
    v_res = client.post("/api/ventas/", json={
        "items": [{"paca_categoria_id": cat_id, "cantidad": 2}]
    })
    assert v_res.status_code == 201
    venta_total = v_res.json()["total"]
    assert venta_total == 600.0

    # Fetch dashboard
    dash_res = client.get("/api/reportes/dashboard", headers=admin_headers)
    assert dash_res.status_code == 200
    data = dash_res.json()

    assert data["ventas_hoy"] >= venta_total
    assert data["ventas_semana"] >= data["ventas_hoy"]
    assert data["inversion_total"] == 2000.0
    assert data["total_prendas"] == 10
    assert data["prendas_vendidas"] == 2
    assert data["prendas_disponibles"] == 8
    assert data["ganancia_neta"] == 600.0 - 2000.0


def test_excel_export_and_formula_injection(client, admin_headers):
    # Seed a paca with formula injection payload in categoria name
    injection_name = '=HYPERLINK("https://evil.example","Hacked")'
    p_res = client.post("/api/pacas/", json={
        "descripcion": "Paca Security Excel",
        "costo": 500.0,
        "peso_lbs": 10.0,
        "categorias": [
            {"nombre": injection_name, "cantidad_total": 5, "precio_venta": 100.0}
        ]
    }, headers=admin_headers)
    cat_id = p_res.json()["categorias"][0]["id"]

    # Register a sale with this category
    client.post("/api/ventas/", json={
        "items": [{"paca_categoria_id": cat_id, "cantidad": 1}]
    })

    # Download Excel
    res = client.get("/api/reportes/excel", headers=admin_headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # Verify workbook content and sanitization using openpyxl
    wb = openpyxl.load_workbook(BytesIO(res.content))
    assert "Ventas" in wb.sheetnames
    assert "Resumen por Paca" in wb.sheetnames

    ws = wb["Ventas"]
    headers = [cell.value for cell in ws[1]]
    assert headers == ["ID Venta", "Fecha", "Categoría", "Cantidad", "Precio Unitario", "Subtotal", "Total Venta"]

    # Verify category cell starts with single quote escaping formula
    cat_cell_val = str(ws.cell(row=2, column=3).value)
    assert cat_cell_val.startswith("'=")
    assert "evil.example" in cat_cell_val
